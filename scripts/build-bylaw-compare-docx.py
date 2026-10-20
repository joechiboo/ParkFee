# 對照表 Word 版（直式 A4）——由 build-bylaw-compare.mjs 呼叫，不單獨維護內容。
#
# 輸入：node 端吐出的列 JSON（每列 new/old/note 三格，HTML 片段只含
#       <u class="r">紅字底線</u>、<br>、<div class="hd[ r]">條名</div> 三種標記）
# 輸出：.docx（表頭跨頁重複、列不跨頁拆開、紅字底線保留，給經理貼進會議紀錄）
#
# 用法：python scripts/build-bylaw-compare-docx.py <rows.json> <out.docx>
import json
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = '新細明體'
RED = RGBColor(0xDD, 0x00, 0x00)
TITLE = '樂菲莊園社區地下室停車場管理辦法修正草案條文對照表'
WIDTHS = [Cm(6.6), Cm(6.6), Cm(4.0)]  # A4 直式版心 17.2cm（左右邊界 1.9cm）


def set_font(run, size):
    run.font.name = FONT
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)


def fill(cell, frag, size):
    """把 HTML 片段寫進儲存格：<div>／<br> 斷段落，<u class="r"> 轉紅字底線。"""
    frag = re.sub(r'<div class="hd( r)?">(.*?)</div>', lambda m: f'<u class="r">{m[2]}</u><br>' if m[1] else f'{m[2]}<br>', frag)
    lines = re.split(r'<br\s*/?>', frag)
    cell.paragraphs[0].text = ''
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        for red, text in _segments(line):
            run = p.add_run(text)
            set_font(run, size)
            if red:
                run.font.color.rgb = RED
                run.font.underline = True


def _segments(line):
    out, pos = [], 0
    for m in re.finditer(r'<u class="r">(.*?)</u>', line):
        if m.start() > pos:
            out.append((False, line[pos:m.start()]))
        out.append((True, m[1]))
        pos = m.end()
    if pos < len(line):
        out.append((False, line[pos:]))
    return [(r, re.sub(r'<[^>]+>', '', t)) for r, t in out if t]


def row_opts(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement('w:cantSplit')
    trPr.append(cant)
    if header:
        h = OxmlElement('w:tblHeader')
        trPr.append(h)


def main(src, out):
    rows = json.load(open(src, encoding='utf-8'))
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(8)
    set_font(t.add_run(TITLE), 15)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, label in enumerate(['修正條文', '原有條文', '修正說明']):
        c = table.rows[0].cells[i]
        c.width = WIDTHS[i]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(label), 11)
    row_opts(table.rows[0], header=True)

    for r in rows:
        cells = table.add_row().cells
        for i, key in enumerate(['new', 'old', 'note']):
            cells[i].width = WIDTHS[i]
            fill(cells[i], r[key], 9.5 if key == 'note' else 10)
        row_opts(table.rows[-1])

    doc.save(out)
    print(f'✓ {out}（{len(rows)} 列）')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
